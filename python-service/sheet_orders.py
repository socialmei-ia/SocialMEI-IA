"""Atendimento experimental por planilha, separado do catálogo compartilhado."""

import argparse
import difflib
import hmac
import os
from pathlib import Path
import re
import tempfile
import threading
import time
import unicodedata
from dataclasses import dataclass
from datetime import datetime
from decimal import Decimal, InvalidOperation
from zipfile import BadZipFile
from xml.etree.ElementTree import ParseError

from fastapi import APIRouter, Depends, Header, HTTPException
from openpyxl import Workbook, load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from pydantic import AliasChoices, BaseModel, Field, StrictStr, field_validator

COLUMN_ALIASES = {
    "product": {"produto", "item", "nome", "descricao"},
    "stock": {"estoque", "quantidade", "qtd", "qtde", "estoque_atual", "saldo"},
    "minimum": {"minimo", "estoque_minimo", "min"},
    "price": {"preco", "valor", "preco_unitario", "valor_unitario"},
}
NUMBER_WORDS = dict(
    zip(
        (
            "um",
            "uma",
            "dois",
            "duas",
            "tres",
            "quatro",
            "cinco",
            "seis",
            "sete",
            "oito",
            "nove",
            "dez",
        ),
        (1, 1, 2, 2, 3, 4, 5, 6, 7, 8, 9, 10),
    )
)
IGNORED_WORDS = set(
    "quero queria gostaria preciso vou querer pedir pedido me ve manda mande de do da dos das o a os as um uma por favor pra para oi ola bom dia boa tarde noite".split()
)
CONFIRMATIONS = {"confirmar", "confirmo", "sim", "fechar", "fechado"}
CANCELLATIONS = {"cancelar", "cancela", "nao"}
HELP = "Oi! Me diga o que você quer, por exemplo: *2 canecas e 1 chaveiro*.\nPara ver o estoque, digite *estoque*."
MAX_PENDING = 500
PENDING_TTL = 30 * 60


class SpreadsheetError(Exception):
    pass


def normalize(text):
    value = (
        unicodedata.normalize("NFKD", str(text))
        .encode("ascii", "ignore")
        .decode()
        .lower()
    )
    return re.sub(r"[^\w\s,;+]", " ", value).strip()


def parse_number(value, *, optional=False, integer=False):
    if value is None or str(value).strip() == "":
        if optional:
            return None
        raise SpreadsheetError(
            "Há um estoque vazio. Preencha quantidades inteiras não negativas."
        )
    try:
        if isinstance(value, bool):
            raise ValueError
        text = str(value).strip()
        if "," in text:
            text = text.replace(".", "").replace(",", ".")
        number = Decimal(text)
        if (
            not number.is_finite()
            or number < 0
            or (integer and number != number.to_integral_value())
        ):
            raise ValueError
        return int(number) if integer else number
    except (InvalidOperation, ValueError):
        raise SpreadsheetError(
            "Há um valor inválido na planilha. Confira estoque, mínimo e preço."
        ) from None


def currency(value):
    return "R$ " + f"{value:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


@dataclass
class Product:
    row: int
    name: str
    stock: int
    minimum: int | None
    price: Decimal | None

    def status(self):
        if self.minimum is None:
            return ""
        if self.stock < self.minimum / 2:
            return " - CRÍTICO"
        return " - repor" if self.stock < self.minimum else ""

    def public(self):
        return {
            "produto": self.name,
            "estoque": self.stock,
            "minimo": self.minimum,
            "preco": float(self.price) if self.price is not None else None,
        }


def parse_items(text):
    if re.search(r"(?<!\w)[-−]\s*\d|\b\d+[.,]\d+\b", text):
        raise ValueError(
            "Use quantidades inteiras positivas, por exemplo: *2 canecas e 1 chaveiro*."
        )
    items = []
    for part in re.split(r"\s*(?:,|;|\+|\be\b|\bmais\b)\s*", normalize(text)):
        quantity, name = None, []
        for token in part.split():
            match = re.fullmatch(r"x?(\d+)x?", token)
            if quantity is None and match:
                quantity = int(match[1])
            elif quantity is None and token in NUMBER_WORDS:
                quantity = NUMBER_WORDS[token]
            elif token not in IGNORED_WORDS:
                name.append(token)
        if name:
            items.append((" ".join(name), 1 if quantity is None else quantity))
    return items


def tokens(text):
    return {
        word[:-1] if len(word) > 3 and word.endswith("s") else word
        for word in normalize(text).split()
    }


def find_product(products, query):
    normalized = normalize(query)
    exact = [product for product in products if normalize(product.name) == normalized]
    candidates = exact or [
        product
        for product in products
        if tokens(query) and tokens(query) <= tokens(product.name)
    ]
    if len(candidates) == 1:
        return candidates[0], "ok", []
    if candidates:
        return None, "ambiguous", [product.name for product in candidates]
    names = {normalize(product.name): product.name for product in products}
    suggestions = [
        names[name]
        for name in difflib.get_close_matches(normalized, list(names), n=3, cutoff=0.5)
    ]
    return None, "missing", suggestions


class SpreadsheetOrders:
    def __init__(self, path):
        self.path = Path(path)
        self.lock = threading.RLock()
        self.pending = {}

    def open(self):
        if not self.path.is_file():
            raise SpreadsheetError(
                "Planilha não encontrada. O responsável precisa configurar o arquivo de estoque."
            )
        try:
            workbook = load_workbook(self.path)
        except (OSError, ValueError, BadZipFile, InvalidFileException, ParseError):
            raise SpreadsheetError(
                "Não consegui abrir a planilha de estoque."
            ) from None
        try:
            sheet = workbook.worksheets[0]
            columns = {}
            for cell in sheet[1]:
                title = (
                    normalize(cell.value).replace(" ", "_")
                    if cell.value is not None
                    else ""
                )
                for field, aliases in COLUMN_ALIASES.items():
                    if title in aliases and field not in columns:
                        columns[field] = cell.column
            if not {"product", "stock"} <= columns.keys():
                raise SpreadsheetError(
                    "A primeira linha precisa das colunas Produto e Estoque."
                )
            products, names = [], set()
            for row in range(2, sheet.max_row + 1):
                name = sheet.cell(row, columns["product"]).value
                if name is None or not str(name).strip():
                    continue
                name = str(name).strip()
                if normalize(name) in names:
                    raise SpreadsheetError(
                        "Há nomes de produtos duplicados na planilha. Use nomes únicos."
                    )
                names.add(normalize(name))
                products.append(
                    Product(
                        row,
                        name,
                        parse_number(
                            sheet.cell(row, columns["stock"]).value, integer=True
                        ),
                        (
                            parse_number(
                                sheet.cell(row, columns["minimum"]).value,
                                optional=True,
                                integer=True,
                            )
                            if "minimum" in columns
                            else None
                        ),
                        (
                            parse_number(
                                sheet.cell(row, columns["price"]).value, optional=True
                            )
                            if "price" in columns
                            else None
                        ),
                    )
                )
            return workbook, sheet, columns, products
        except Exception:
            workbook.close()
            raise

    def save(self, workbook):
        temporary = None
        try:
            with tempfile.NamedTemporaryFile(
                dir=self.path.parent, suffix=".xlsx", delete=False
            ) as file:
                temporary = Path(file.name)
            workbook.save(temporary)
            os.replace(temporary, self.path)
        except OSError:
            raise SpreadsheetError(
                "Não consegui salvar o estoque. O responsável precisa conferir o arquivo e as permissões."
            ) from None
        finally:
            if temporary is not None:
                temporary.unlink(missing_ok=True)

    def stock(self):
        with self.lock:
            workbook, _, _, products = self.open()
            try:
                return [product.public() for product in products]
            finally:
                workbook.close()

    def new_order(self, session, items, products):
        problems, accepted = [], {}
        for name, quantity in items:
            if not 0 < quantity <= 1_000_000:
                problems.append(f'- "{name}": quantidade inválida')
                continue
            product, status, suggestions = find_product(products, name)
            if status == "ambiguous":
                problems.append(f'- "{name}": qual destes? ' + ", ".join(suggestions))
            elif status == "missing":
                hint = f" Quis dizer: {', '.join(suggestions)}?" if suggestions else ""
                problems.append(f'- "{name}": não encontrei esse produto.{hint}')
            else:
                accepted[product.name] = accepted.get(product.name, 0) + quantity
        by_name = {product.name: product for product in products}
        for name, quantity in accepted.items():
            if quantity > by_name[name].stock:
                problems.append(
                    f"- {name}: só temos {by_name[name].stock} (você pediu {quantity})"
                )
        if problems:
            self.pending.pop(session, None)
            return (
                "Não consegui fechar o pedido:\n"
                + "\n".join(problems)
                + "\n\nQuer ajustar?"
            )
        if session not in self.pending and len(self.pending) >= MAX_PENDING:
            return (
                "Há muitos pedidos em preparação. Tente novamente em alguns instantes."
            )
        self.pending[session] = (accepted, time.monotonic() + PENDING_TTL)
        lines = [f"- {quantity}x {name}" for name, quantity in accepted.items()]
        total = ""
        if all(by_name[name].price is not None for name in accepted):
            total = "\nTotal: " + currency(
                sum(
                    (
                        by_name[name].price * quantity
                        for name, quantity in accepted.items()
                    ),
                    Decimal(0),
                )
            )
        return (
            "Seu pedido:\n"
            + "\n".join(lines)
            + total
            + "\n\nResponda *CONFIRMAR* para fechar ou *CANCELAR*."
        )

    def confirm(self, session, workbook, sheet, columns, products):
        pending = self.pending.get(session)
        if not pending:
            return "Não tenho nenhum pedido aguardando confirmação. " + HELP, None
        order = pending[0]
        by_name = {product.name: product for product in products}
        for name, quantity in order.items():
            if name not in by_name or quantity > by_name[name].stock:
                self.pending.pop(session, None)
                return (
                    f"O estoque de *{name}* mudou e não dá mais para fechar. Faça o pedido de novo.",
                    None,
                )
        if "Pedidos" not in workbook.sheetnames:
            workbook.create_sheet("Pedidos").append(
                ["data", "sessao", "produto", "qtd"]
            )
        history = workbook["Pedidos"]
        warnings = []
        for name, quantity in order.items():
            product = by_name[name]
            product.stock -= quantity
            sheet.cell(product.row, columns["stock"]).value = product.stock
            history.append(
                [datetime.now().strftime("%d/%m/%Y %H:%M"), session, name, quantity]
            )
            # Mantém identificadores externos como texto, nunca como fórmula do Excel.
            for column in (1, 2, 3):
                history.cell(history.max_row, column).data_type = "s"
            if product.status():
                warnings.append(
                    f"- {name}: sobraram {product.stock} (mínimo {product.minimum})"
                )
        self.save(workbook)
        self.pending.pop(session, None)
        return "Pedido confirmado! Obrigado.", (
            "Estoque baixo:\n" + "\n".join(warnings) if warnings else None
        )

    def handle(self, message, session):
        with self.lock:
            now = time.monotonic()
            self.pending = {
                key: value for key, value in self.pending.items() if value[1] > now
            }
            normalized = " ".join(normalize(message).split())
            words = set(normalized.split())
            alert = None
            if normalized in CANCELLATIONS:
                response = (
                    "Pedido cancelado. Se quiser, é só pedir de novo."
                    if self.pending.pop(session, None)
                    else "Não tenho nenhum pedido aguardando confirmação. " + HELP
                )
            elif (
                words & (CONFIRMATIONS | CANCELLATIONS)
                and normalized not in CONFIRMATIONS
            ):
                response = "Para fechar, responda apenas *CONFIRMAR*. Para desistir, responda *CANCELAR*."
            else:
                workbook, sheet, columns, products = self.open()
                try:
                    if normalized in CONFIRMATIONS:
                        response, alert = self.confirm(
                            session, workbook, sheet, columns, products
                        )
                    elif words & {"estoque", "relatorio", "saldo"} and not re.search(
                        r"\d", normalized
                    ):
                        response = (
                            "*Estoque atual*\n"
                            + "\n".join(
                                f"- {product.name}: {product.stock}{product.status()}"
                                for product in products
                            )
                            if products
                            else "A planilha está sem produtos."
                        )
                    else:
                        try:
                            items = parse_items(message)
                            response = (
                                self.new_order(session, items, products)
                                if items
                                else HELP
                            )
                        except ValueError as error:
                            response = str(error)
                finally:
                    workbook.close()
            return {
                "resposta": response,
                "alerta_estoque": alert,
                "pedido_pendente": session in self.pending,
            }


class MessageInput(BaseModel):
    message: StrictStr = Field(
        validation_alias=AliasChoices("mensagem", "chatInput"),
        min_length=1,
        max_length=2000,
    )
    session: StrictStr = Field(
        validation_alias=AliasChoices("sessionId", "sessao"),
        min_length=1,
        max_length=128,
    )

    @field_validator("message", "session")
    @classmethod
    def nonempty(cls, value):
        value = value.strip()
        if not value:
            raise ValueError("Informe uma mensagem e um identificador de sessão.")
        return value


def require_api_key(x_api_key: str | None = Header(default=None, alias="X-API-Key")):
    expected = os.getenv("API_KEY", "")
    if not expected:
        raise HTTPException(503, "API de planilha não configurada.")
    if not hmac.compare_digest((x_api_key or "").encode(), expected.encode()):
        raise HTTPException(401, "Não autorizado.")


def create_router(service=None):
    if os.getenv("SOCIALMEI_SHARED_DATA_ENABLED", "false").lower() in {
        "true",
        "1",
        "yes",
        "on",
    }:
        raise RuntimeError(
            "A API experimental de planilha não pode usar a instância compartilhada do CRM."
        )
    if os.getenv("WEB_CONCURRENCY", "1") != "1":
        raise RuntimeError("A API experimental de planilha exige um único worker.")
    service = service or SpreadsheetOrders(
        os.getenv("PLANILHA", "estoque_exemplo.xlsx")
    )
    router = APIRouter(tags=["spreadsheet-prototype"])

    @router.get("/saude")
    def health():
        return {"status": "ok"}

    @router.get("/estoque", dependencies=[Depends(require_api_key)])
    def stock():
        try:
            return {"produtos": service.stock()}
        except SpreadsheetError as error:
            raise HTTPException(503, str(error)) from None

    @router.post("/mensagem", dependencies=[Depends(require_api_key)])
    def message(payload: MessageInput):
        try:
            return service.handle(payload.message, payload.session)
        except SpreadsheetError as error:
            raise HTTPException(503, str(error)) from None

    return router


def create_example(path):
    path = Path(path)
    if path.exists():
        raise SpreadsheetError("O arquivo já existe. A planilha não foi sobrescrita.")
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Estoque"
    sheet.append(["Produto", "Estoque", "Mínimo", "Preço"])
    for row in [
        ("Camiseta básica P", 25, 15, 39.9),
        ("Camiseta básica M", 2, 15, 39.9),
        ("Bolsa de tecido", 22, 10, 29.9),
        ("Caneca personalizada", 8, 10, 24.9),
        ("Chaveiro", 90, 30, 9.9),
    ]:
        sheet.append(row)
    for letter, width in zip("ABCD", (26, 10, 10, 10)):
        sheet.column_dimensions[letter].width = width
    try:
        # Criação exclusiva: não substitui um arquivo surgido entre a conferência e o save.
        with path.open("xb") as target:
            workbook.save(target)
    finally:
        workbook.close()


def cli():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--example", "--exemplo", type=Path)
    parser.add_argument("--chat", "--conversar", action="store_true")
    args = parser.parse_args()
    try:
        if args.example:
            create_example(args.example)
            print("Planilha fictícia criada. Nenhum estoque real foi alterado.")
        elif args.chat:
            service = SpreadsheetOrders(os.getenv("PLANILHA", "estoque_exemplo.xlsx"))
            while message := input("você: ").strip():
                result = service.handle(message, "terminal")
                print("bot:", result["resposta"])
                if result["alerta_estoque"]:
                    print("[Aviso interno]", result["alerta_estoque"])
        else:
            parser.print_help()
    except (SpreadsheetError, OSError) as error:
        print(error)
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(cli())
