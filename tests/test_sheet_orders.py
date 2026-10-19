"""Testes isolados: usam apenas planilhas temporárias com dados fictícios."""

import hashlib
import os
from pathlib import Path
import sys
import tempfile
import subprocess
import unittest
from concurrent.futures import ThreadPoolExecutor
from unittest.mock import patch

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "python-service"))

from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from sheet_orders import (
    SpreadsheetError,
    SpreadsheetOrders,
    create_example,
    create_router,
)


class SpreadsheetOrdersTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.path = Path(self.directory.name) / "example.xlsx"
        create_example(self.path)
        self.service = SpreadsheetOrders(self.path)

    def stock(self, name="Caneca personalizada"):
        return next(
            item["estoque"] for item in self.service.stock() if item["produto"] == name
        )

    def test_order_only_changes_stock_after_confirmation(self):
        before = self.path.read_bytes()
        order = self.service.handle("quero 2 canecas e 1 chaveiro", "alice")
        self.assertTrue(order["pedido_pendente"])
        self.assertIn("R$ 59,70", order["resposta"])
        self.assertEqual(before, self.path.read_bytes())
        result = self.service.handle("CONFIRMAR", "alice")
        self.assertEqual("Pedido confirmado! Obrigado.", result["resposta"])
        self.assertIn("sobraram 6", result["alerta_estoque"])
        self.assertFalse(result["pedido_pendente"])
        self.assertEqual(6, self.stock())
        self.assertEqual(89, self.stock("Chaveiro"))
        workbook = load_workbook(self.path)
        self.addCleanup(workbook.close)
        self.assertEqual(3, workbook["Pedidos"].max_row)
        self.service.handle("CONFIRMAR", "alice")
        self.assertEqual(6, self.stock())

    def test_cancellation_and_negated_confirmation_do_not_decrement(self):
        self.service.handle("duas canecas", "alice")
        result = self.service.handle("não confirmar", "alice")
        self.assertTrue(result["pedido_pendente"])
        self.assertEqual(8, self.stock())
        result = self.service.handle("CANCELAR", "alice")
        self.assertFalse(result["pedido_pendente"])
        self.service.handle("CONFIRMAR", "alice")
        self.assertEqual(8, self.stock())

    def test_sessions_recheck_available_stock_on_confirmation(self):
        self.service.handle("5 canecas", "alice")
        self.service.handle("5 canecas", "bob")
        with ThreadPoolExecutor(max_workers=2) as executor:
            results = list(
                executor.map(
                    lambda session: self.service.handle("confirmar", session),
                    ["alice", "bob"],
                )
            )
        self.assertEqual(3, self.stock())
        self.assertEqual(
            1,
            sum(
                result["resposta"] == "Pedido confirmado! Obrigado."
                for result in results
            ),
        )
        self.assertTrue(any("estoque" in result["resposta"] for result in results))

    def test_aggregated_and_invalid_quantities_are_rejected(self):
        for message in (
            "2 canecas e 7 canecas",
            "0 canecas",
            "-2 canecas",
            "1,5 canecas",
            "1000001 canecas",
        ):
            with self.subTest(message=message):
                self.assertFalse(
                    self.service.handle(message, "alice")["pedido_pendente"]
                )
                self.assertEqual(8, self.stock())

    def test_ambiguity_and_spelling_suggestions_never_choose_a_product(self):
        for message in ("2 camisetas", "2 canec"):
            with self.subTest(message=message):
                self.assertFalse(
                    self.service.handle(message, "alice")["pedido_pendente"]
                )
        self.assertEqual(8, self.stock())

    def test_expired_order_cannot_be_confirmed(self):
        with patch("sheet_orders.time.monotonic", return_value=100):
            self.service.handle("2 canecas", "alice")
        with patch("sheet_orders.time.monotonic", return_value=2000):
            result = self.service.handle("confirmar", "alice")
        self.assertFalse(result["pedido_pendente"])
        self.assertEqual(8, self.stock())

    def test_failed_replace_keeps_original_and_pending_order(self):
        self.service.handle("2 canecas", "alice")
        original = hashlib.sha256(self.path.read_bytes()).digest()
        with patch("sheet_orders.os.replace", side_effect=PermissionError):
            with self.assertRaises(SpreadsheetError):
                self.service.handle("confirmar", "alice")
        self.assertEqual(original, hashlib.sha256(self.path.read_bytes()).digest())
        self.assertIn("alice", self.service.pending)
        self.assertEqual([self.path], list(self.path.parent.iterdir()))

    def test_external_session_is_written_as_text_in_excel(self):
        self.service.handle("1 chaveiro", '=HYPERLINK("example")')
        self.service.handle("confirmar", '=HYPERLINK("example")')
        workbook = load_workbook(self.path)
        self.addCleanup(workbook.close)
        self.assertEqual("s", workbook["Pedidos"]["B2"].data_type)

    def test_example_never_overwrites_existing_file(self):
        original = self.path.read_bytes()
        with self.assertRaises(SpreadsheetError):
            create_example(self.path)
        self.assertEqual(original, self.path.read_bytes())

    def test_optional_columns_and_portuguese_numbers(self):
        workbook = load_workbook(self.path)
        sheet = workbook.active
        sheet.delete_cols(3, 2)
        workbook.save(self.path)
        workbook.close()
        result = self.service.handle("quero duas canecas", "alice")
        self.assertNotIn("Total:", result["resposta"])
        result = self.service.handle("confirmar", "alice")
        self.assertIsNone(result["alerta_estoque"])
        self.assertEqual(6, self.stock())

    def test_invalid_workbooks_are_rejected(self):
        for invalid in (-1, "NaN", 1.5, "=1+1"):
            with self.subTest(invalid=invalid):
                workbook = load_workbook(self.path)
                workbook.active["B2"] = invalid
                workbook.save(self.path)
                workbook.close()
                with self.assertRaises(SpreadsheetError):
                    self.service.stock()

    def test_duplicate_product_names_are_rejected(self):
        workbook = load_workbook(self.path)
        workbook.active["A3"] = "CAMISETA BÁSICA P"
        workbook.save(self.path)
        workbook.close()
        with self.assertRaises(SpreadsheetError):
            self.service.stock()


class SpreadsheetHTTPTests(unittest.TestCase):
    def setUp(self):
        self.environment = patch.dict(
            os.environ,
            {
                "SOCIALMEI_SHARED_DATA_ENABLED": "false",
                "WEB_CONCURRENCY": "1",
                "API_KEY": "test-only-key",
            },
        )
        self.environment.start()
        self.addCleanup(self.environment.stop)
        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        path = Path(self.directory.name) / "example.xlsx"
        create_example(path)
        self.service = SpreadsheetOrders(path)
        app = FastAPI()
        app.include_router(create_router(self.service))
        self.client = TestClient(app)
        self.addCleanup(self.client.close)
        self.headers = {"X-API-Key": "test-only-key"}

    def test_health_and_key_protection(self):
        self.assertEqual({"status": "ok"}, self.client.get("/saude").json())
        self.assertEqual(401, self.client.get("/estoque").status_code)
        self.assertEqual(
            401,
            self.client.post(
                "/mensagem", json={"mensagem": "estoque", "sessionId": "a"}
            ).status_code,
        )
        self.assertEqual(
            5, len(self.client.get("/estoque", headers=self.headers).json()["produtos"])
        )
        with patch.dict(os.environ, {"API_KEY": ""}):
            self.assertEqual(
                503, self.client.get("/estoque", headers=self.headers).status_code
            )

    def test_n8n_aliases_and_required_session(self):
        result = self.client.post(
            "/mensagem",
            headers=self.headers,
            json={"chatInput": "2 canecas", "sessionId": "n8n-session"},
        )
        self.assertEqual(200, result.status_code)
        self.assertTrue(result.json()["pedido_pendente"])
        for payload in (
            {"mensagem": "2 canecas"},
            {"mensagem": 123, "sessionId": "a"},
            {"mensagem": " ", "sessao": "a"},
            {"mensagem": "estoque", "sessao": "a" * 129},
        ):
            self.assertEqual(
                422,
                self.client.post(
                    "/mensagem", headers=self.headers, json=payload
                ).status_code,
            )

    def test_missing_workbook_returns_sanitized_unavailability(self):
        self.service.path.unlink()
        result = self.client.get("/estoque", headers=self.headers)
        self.assertEqual(503, result.status_code)
        self.assertNotIn(self.directory.name, result.text)

    def test_application_opt_in_preserves_existing_endpoints(self):
        script = """
from fastapi.testclient import TestClient
from app import app
with TestClient(app) as client:
    assert client.get('/health').json() == {'status': 'healthy'}
    assert client.post('/processar', json={'texto': 'Olá'}).json()['maiusculo'] == 'OLÁ'
    assert client.get('/saude').status_code == EXPECTED
"""
        for flag, expected in (("false", 404), ("true", 200)):
            environment = dict(os.environ, SOCIALMEI_SHEET_ORDERS_ENABLED=flag)
            result = subprocess.run(
                [sys.executable, "-c", script.replace("EXPECTED", str(expected))],
                cwd=Path(__file__).resolve().parents[1] / "python-service",
                env=environment,
                capture_output=True,
                text=True,
            )
            self.assertEqual(0, result.returncode, result.stderr)

    def test_shared_crm_and_multiple_workers_are_rejected(self):
        for configuration in (
            {"SOCIALMEI_SHARED_DATA_ENABLED": "true"},
            {"WEB_CONCURRENCY": "2"},
        ):
            with self.subTest(configuration=configuration), patch.dict(
                os.environ, configuration
            ):
                with self.assertRaises(RuntimeError):
                    create_router(self.service)


if __name__ == "__main__":
    unittest.main()
